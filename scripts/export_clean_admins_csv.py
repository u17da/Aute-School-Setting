import openpyxl
import csv
import json

wb = openpyxl.load_workbook('data/all_schools_original.xlsx', data_only=True)
print('Sheet names:', wb.sheetnames)

# 性別マッピングの構築 (Sheet2)
gender_map = {}
if 'Sheet2' in wb.sheetnames:
    ws_gender = wb['Sheet2']
    for r in range(2, ws_gender.max_row + 1):
        uid = str(ws_gender.cell(r, 1).value or '').strip()
        g = str(ws_gender.cell(r, 2).value or '').strip().lower()
        if uid and g:
            gender_map[uid] = g
print(f'性別マッピング数: {len(gender_map)} 件')

# 管理者リストシートの解析
ws = wb['管理者リスト'] if '管理者リスト' in wb.sheetnames else wb.worksheets[0]

admin_records = []
for r in range(2, ws.max_row + 1):
    school_code = str(ws.cell(r, 1).value or '').strip()
    admin_plus_id = str(ws.cell(r, 2).value or '').strip()
    admin_plus_pw = str(ws.cell(r, 3).value or '').strip()
    school_name = str(ws.cell(r, 4).value or '').strip()
    role = str(ws.cell(r, 5).value or 'admin').strip()
    display_name = str(ws.cell(r, 6).value or '').strip()
    gender = str(ws.cell(r, 7).value or '').strip().lower()
    email = str(ws.cell(r, 8).value or '').strip()
    user_id = str(ws.cell(r, 9).value or '').strip()
    password = str(ws.cell(r, 10).value or '').strip()
    federation_id = str(ws.cell(r, 11).value or '').strip()
    job_title = str(ws.cell(r, 15).value or '').strip()

    if not school_code or not user_id:
        continue

    # 性別がVLOOKUPで空または数式文字列の場合はSheet2から補完
    if not gender or 'vlookup' in gender:
        gender = gender_map.get(user_id, 'male')
    if gender not in ['male', 'female']:
        gender = 'male'

    admin_records.append({
        'schoolCode': school_code,
        'adminPlusId': admin_plus_id,
        'adminPlusPw': admin_plus_pw,
        'schoolName': school_name,
        'role': role,
        'displayName': display_name,
        'gender': gender,
        'email': email,
        'userId': user_id,
        'password': password,
        'federationId': federation_id,
        'jobTitle': job_title
    })

print(f'有効な管理者レコード数: {len(admin_records)} 件')

schools = sorted(list(set(r['schoolCode'] for r in admin_records)))
idp_schools = sorted(list(set(r['schoolCode'] for r in admin_records if '@' in r['adminPlusId'])))
local_schools = sorted(list(set(r['schoolCode'] for r in admin_records if '@' not in r['adminPlusId'])))

print(f'総学校数: {len(schools)} 校')
print(f'  - 外部IdP認証校: {len(idp_schools)} 校')
print(f'  - ローカル認証校: {len(local_schools)} 校')

# CSV書き出し (BOM付きUTF-8)
out_csv = 'data/all_school_admins_clean.csv'
header = ['学校コード', 'admin+ ID', 'admin+ PW', '学校名', '役割', '表示名', '性別', 'メールアドレス', 'ユーザーID', 'パスワード', '外部認証ID', '役職']
with open(out_csv, 'w', encoding='utf-8-sig', newline='') as f:
    writer = csv.writer(f)
    writer.writerow(header)
    for r in admin_records:
        writer.writerow([
            r['schoolCode'],
            r['adminPlusId'],
            r['adminPlusPw'],
            r['schoolName'],
            r['role'],
            r['displayName'],
            r['gender'],
            r['email'],
            r['userId'],
            r['password'],
            r['federationId'],
            r['jobTitle']
        ])

print(f'CSV出力完了: {out_csv}')

# サマリーJSON出力
summary = {
    'totalAdmins': len(admin_records),
    'totalSchools': len(schools),
    'idpSchoolsCount': len(idp_schools),
    'localSchoolsCount': len(local_schools),
    'idpSchools': idp_schools,
    'localSchools': local_schools,
    'sampleRecords': admin_records[:5]
}
with open('data/all_school_admins_summary.json', 'w', encoding='utf-8') as f:
    json.dump(summary, f, ensure_ascii=False, indent=2)
print('サマリー出力完了: data/all_school_admins_summary.json')
